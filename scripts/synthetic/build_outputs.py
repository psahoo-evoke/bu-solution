"""Write the Excel workbook, CSV and JSON seed files."""
import json, os, sys, csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1]
D = json.load(open(os.path.join(HERE, "data.json")))
M = json.load(open(os.path.join(HERE, "manifest.json")))
SCH = {s["school_code"]: s for s in D["schools"]}

RULES = [
 ("R01","Accreditation","Only courses from regionally accredited institutions are eligible.","school.accreditation","No credit, TFRGEN NOCRDT, 0 units, TX","NC-ACCRED","Sponsor Q8"),
 ("R02","Course type","Remedial/developmental, vocational, physical education and ROTC courses are ineligible.","course.course_type","No credit, TX","NC-REMEDIAL / NC-VOC / NC-PE / NC-ROTC","Sponsor Q8"),
 ("R03","Course type","Internships, research and military or life experience are ineligible.","course.course_type","No credit, TX","NC-INTERNSHIP / NC-RESEARCH","Sponsor Q8"),
 ("R04","Credits","One-credit courses are generally not accepted.","credits = 1","No credit, TX","NC-ONECREDIT","Sponsor Q8"),
 ("R05","Grades","Pass/Fail grades are generally not accepted.","grade in (P, NP, S, U)","No credit, TX","NC-PASSFAIL","Sponsor Q8"),
 ("R06","Grades","Minimum grade C; C- and below are ineligible. External grades do not enter the BU GPA.","min_grade = C","No credit, TX","NC-GRADE","Brief; sponsor Q18"),
 ("R07","Age","Courses five or more years old need review before any credit.","term_year <= start_year - 5","Pending: age review","RV-AGE","Sponsor Q8"),
 ("R08","Conversion","Quarter credits x 2/3, rounded to 2 decimals, course by course.","calendar = Quarter","units = ROUND(credits*2/3, 2)","-","Sponsor Q18/Q27"),
 ("R09","Units","Units posted are always the units the student took (converted), never the BU course's credits.","-","-","-","Sponsor Q27"),
 ("R10","Existing match","Check the equivalency table first; a reusable decision is applied without a syllabus.","reusable = Yes","Matched: BU course or TR elective, grade T","(decision's code)","Sponsor Q1, Q23"),
 ("R11","Conflicts","Conflicting past decisions: the newest decision for the same catalog years wins; flag the conflict for the department.","multiple decisions","Apply newest; flag","-","Sponsor Q18"),
 ("R12","Duplicates","A second course covering the same BU course earns elective credit only.","bu_course already used","Elective (TR)","EL-DUPLICATE","Sponsor Q8/Q18"),
 ("R13","Residency","CS Group C (CS 320, 332, 350) must be taken at BU: requirement blocked, units still post as elective.","bu_course in Group C","Elective (TR)","EL-RESIDENCY","Sponsor Q29/Q31"),
 ("R14","Residency","Biology: at least five BI courses beyond BI 107/108 must be taken at BU (advisory; credit unaffected).","major = Biology","Advisory flag","-","Sponsor Q29/Q31"),
 ("R15","Residency","Economics: at most four 200+ level courses transferred AFTER matriculation.","term after BU start","Excess -> elective","EL-RESIDENCY","Sponsor Q8/Q31"),
 ("R16","Residency","Engineering: at least 48 upper-division units must be taken at BU (not in this data set's majors).","school = ENG","Advisory flag","-","Sponsor Q29/Q31"),
 ("R17","Cap","At most 68 external units can count toward the degree.","sum(units) > 68","Excess units do not count; flag","CAP-68","Sponsor Q8/Q31"),
 ("R18","Hub","Transfer courses do not earn BU Hub general education units.","-","-","-","Brief"),
 ("R19","Deadline","All credits must be transferred within one year of starting at BU; after that an Outstanding Transfer Credit Review is needed.","today > start + 1 year","Escalate","-","Brief; sponsor Q32"),
 ("R20","Matching tiers","Learning-outcome overlap >=80% -> propose equivalent; 50-79% -> evaluator decides; <50% -> elective if eligible.","AI overlap score","Recommendation only","-","Sponsor Q7"),
 ("R21","Batch approval","Recommendations above 90% may be approved together as a reviewed batch; each course still logged.","AI confidence > 90%","Batch list","-","Sponsor Q11"),
 ("R22","Evidence","Course descriptions are not accepted as evidence; they may drive preliminary AI suggestions only.","no syllabus","Preliminary; not decidable","PD-TS","Sponsor Q16"),
 ("R23","Syllabus chase","Missing syllabus -> TS; incomplete -> TM request. Reminders on days 7, 21 and 45. A current syllabus from another instructor of the same code/title is accepted.","syllabus state","Status + reminders","PD-TS / PD-TM","Sponsor Q16, Q22"),
 ("R24","Codes","Equivalent -> specific BU course; elective -> BU code ending in TR; no credit -> TFRGEN NOCRDT, 0 units, TX. Every evaluated course shows grade T (TX for no credit).","-","-","-","Sponsor Q24/Q27"),
]
REASONS = [
 ("EQ-CONTENT","Equivalent","Substantially equivalent content and level (>=80% overlap)"),
 ("EQ-EVALUATOR","Equivalent","Partial overlap; evaluator judged it equivalent"),
 ("EL-PARTIAL","Elective","Partial overlap; evaluator judged elective"),
 ("EL-LOWOVERLAP","Elective","Below 50% overlap; eligible course -> elective"),
 ("EL-NOEQUIV","Elective","Eligible course with no BU equivalent"),
 ("EL-DUPLICATE","Elective","Second course covering the same BU course"),
 ("EL-RESIDENCY","Elective","Requirement must be taken at BU; units post as elective"),
 ("NC-GRADE","No credit","Grade below C"),("NC-PASSFAIL","No credit","Pass/Fail grade"),
 ("NC-ONECREDIT","No credit","One-credit course"),("NC-REMEDIAL","No credit","Remedial / developmental course"),
 ("NC-PE","No credit","Physical education activity course"),("NC-VOC","No credit","Vocational course"),
 ("NC-ROTC","No credit","ROTC / military science"),("NC-INTERNSHIP","No credit","Internship or co-op"),
 ("NC-RESEARCH","No credit","Research / independent study"),("NC-ACCRED","No credit","Institution not regionally accredited"),
 ("RV-AGE","Pending","Course 5+ years old; needs review"),("PD-TS","Pending","Syllabus requested (TS)"),
 ("PD-TM","Pending","Additional materials requested (TM)"),("PD-REVIEW","Pending","Syllabus on file; awaiting evaluator"),
 ("CAP-68","Flag","68-unit external cap reached"),
]
TESTS = [
 ("T01","Sponsor example","3-credit semester film course, grade B, no BU equivalent","Elective, 3.00 units, grade T"),
 ("T02","Sponsor example","Maya: PSC PSY 215, 4 quarter credits, grade B, no equivalent","Elective CAS PS 2TR, 2.67 units (4 x 2/3), grade T"),
 ("T03","Sponsor example","Jordan: NU MATH 1341 with grade C-","No credit, 0 units, TFRGEN NOCRDT, TX"),
 ("T04","Sponsor example","Jordan: NU MATH 0990 developmental math, grade A","No credit, 0 units, TFRGEN NOCRDT, TX"),
 ("T05","Lab credits","Aisha: NU 5-credit lab courses (BIOL 1107)","5.00 units (student's units, not BU's 4)"),
 ("T06","Duplicate","Jordan: NU CS 1990 after CS 2000 (both CAS CS 111)","Second course -> CAS CS 1TR elective, EL-DUPLICATE"),
 ("T07","Residency","Jordan: NU CS 4400 maps to CAS CS 320 (Group C)","Elective CAS CS 2TR, 4 units, EL-RESIDENCY"),
 ("T08","Pass/Fail","Jordan: NU FILM 1101 graded P","No credit, NC-PASSFAIL"),
 ("T09","Age","Jordan: NU HIST 1130 from Fall 2019, BU start Fall 2026","Pending age review, RV-AGE"),
 ("T10","Accreditation","Liam: two Summit Online Institute courses (fictional, not regionally accredited)","No credit, NC-ACCRED"),
 ("T11","68-unit cap","Liam: total countable units exceed 68","Units beyond 68 do not count; CAP-68 flag on the courses that cross it"),
 ("T12","Conflict (superseded)","Liam: NYU ECON-UA 18 has 2022-23 elective and 2024-25 equivalent decisions","Apply newest (CAS EC 203); note superseded decision"),
 ("T13","Conflict (same year)","Aisha: NU BIOL 2703 Microbiology has two different decisions in 2024-25","Apply newest; flag conflict for the department"),
 ("T14","Econ cap timing","Liam: Economics 200+ courses taken before matriculation","Cap not applied (applies only after matriculation)"),
 ("T15","Reuse after approval","Maya's NYU PSYCH-UA 10 is approved in the demo; Liam took the same course","After approval, Liam's line matches instantly"),
 ("T16","Incomplete syllabus","Jordan: NU CS 2700 syllabus has no outcomes or schedule","TM request; no AI recommendation"),
 ("T17","Prompt injection","Liam: NYU ECON-UA 266 syllabus contains 'ignore all previous instructions...'","Injection ignored; tier stays 'evaluator decides' (~62%)"),
 ("T18","AI quality gate","30 held-out decisions (Past Decisions, split = holdout)","Top-1 agreement >= 80%; wrongly suggested credit <= 2%"),
 ("T19","Split course","Jordan: NU CS 2100 (~62%) and CS 2700 both relate to CAS CS 112 (workbook note)","Evaluator may need both syllabi; AI shows the relationship"),
 ("T20","Combined course","Aisha: NU CHEM 1161 one-semester general chemistry","Evaluator decides CH 101 vs CH 102 (workbook note)"),
]

HDR_FILL = PatternFill("solid", fgColor="1F3A5F"); HDR_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
GT_FILL = PatternFill("solid", fgColor="7B2D26")
BODY = Font(name="Arial", size=10); BOLD = Font(name="Arial", size=10, bold=True)

def sheet(wb, name, headers, rows, widths=None, gt_cols=()):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for i, h in enumerate(headers, 1):
        c = ws.cell(1, i); c.font = HDR_FONT; c.fill = GT_FILL if h in gt_cols else HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY; c.alignment = Alignment(vertical="top", wrap_text=isinstance(c.value, str) and len(c.value) > 40)
    for i, h in enumerate(headers, 1):
        w = (widths or {}).get(h) or min(60, max(10, len(h) + 2, *(len(str(r[i - 1])) * 0.9 for r in rows[:50] if r[i - 1] is not None)))
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    return ws

J = lambda xs: " | ".join(x for x in xs if x)
wb = Workbook(); wb.remove(wb.active)

# README
ws = wb.create_sheet("README")
readme = [
 ("TransferIQ synthetic data set (v2: Northeastern + NYU)", ""),
 ("Purpose", "Test data for the EvoHack 2026 transfer credit evaluation prototype (Boston University problem statement)."),
 ("Status", "TEST DATA. Students, grades, transcripts, syllabus content, instructors and decisions are invented. Transcripts and syllabi were NOT issued by Northeastern or NYU."),
 ("Source schools", "Northeastern University (NEU) and New York University (NYU), as given in the hackathon workbook. Two fictional helper schools cover cases NU/NYU cannot: Pinecrest State College (quarter calendar, 1 course) and Summit Online Institute (not regionally accredited, 2 courses)."),
 ("BU courses", "150 courses across 5 majors. 41 codes/titles come from the provided BU workbook; 109 are synthetic. All descriptions, topics and outcomes are synthetic."),
 ("External courses", "194 courses: 76 real NU/NYU catalog courses (code, title, credits from the workbook; content synthetic) + 115 synthetic NU/NYU-style courses + 3 helper-school courses. See the Source column."),
 ("Ground truth", "Columns with a dark red header (gt_*) are answer keys. For the 76 real courses the counterpart follows the workbook's Course Match sheet. Hide gt_* from the AI matcher."),
 ("Past decisions", "100 decisions = the stand-in equivalency table. split = train (reusable, 70) or holdout (hidden from the app, used to measure AI accuracy, 30)."),
 ("Transcripts", "4 students, 7 transcript PDFs, 56 course lines. Transcript Lines shows the expected rules-engine outcome for each line."),
 ("Syllabi", "12 syllabus PDFs: demo maybe-courses, one incomplete (TM), one with a prompt-injection line (red team), three for held-out courses."),
 ("Codes (assumption)", "Elective codes follow the sponsor's 'ends in TR' rule: CAS <DEPT> 1TR (100-level) / 2TR (200+), TFRGEN 1TR for general electives. No credit = TFRGEN NOCRDT, 0 units, TX."),
 ("Formulas", "Transcript Lines 'Units if approved' and the Students summary columns are live formulas."),
 ("Seed files", "seed/*.csv (one per sheet; list fields joined with ' | ') and seed/transferiq_seed.json (lists kept as arrays)."),
]
for r in readme: ws.append(r)
ws["A1"].font = Font(name="Arial", size=14, bold=True)
for row in ws.iter_rows(min_row=2):
    row[0].font = BOLD; row[1].font = BODY; row[1].alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 20; ws.column_dimensions["B"].width = 120
ws["B3"].font = Font(name="Arial", size=10, bold=True, color="B00020")

# Demo scenario
demo = [
 (1,"Upload","Maya uploads Transcript_SYN-24001_NYU and Transcript_SYN-24001_PSC","6 lines extracted; PSY 215 flagged quarter-system"),
 (2,"Instant matches","NYU PSYCH-UA 1, MATH-UA 121, ECON-UA 2 hit the equivalency table","CAS PS 101, CAS MA 123, CAS EC 101; 4 units each; grade T"),
 (3,"Quarter course","PSC PSY 215 Psychology of Music, 4 qtr credits, B","Elective CAS PS 2TR, 2.67 units"),
 (4,"Maybe 1","NYU PSYCH-UA 10 with Syllabus_NYU_PSYCH-UA-10.pdf on file","AI: ~75% overlap with CAS PS 211 -> evaluator decides (workbook: partial, statistics vs experimental design)"),
 (5,"Maybe 2","NYU PSYCH-UA 25 Cognitive Neuroscience, no syllabus","TS status; preliminary suggestion from description; reminders day 7/21/45"),
 (6,"Evaluator","Open PSYCH-UA 10; comparison already done; approve with reason code","Decision saved, reusable = Yes"),
 (7,"Late upload","Maya uploads Syllabus_NYU_PSYCH-UA-25.pdf","AI: ~62% overlap with CAS PS 231 -> evaluator decides (expected: elective CAS PS 2TR)"),
 (8,"Posting","Department -> Registrar -> MyBU-style CSV","Secured credit updates; Liam's NYU PSYCH-UA 10 now matches instantly"),
]
sheet(wb, "Demo Scenario", ["Step","Stage","Action / input","Expected result"], demo, {"Action / input": 60, "Expected result": 70})

# Students with formulas
ws_s = sheet(wb, "Students", ["Student ID","Name","Intended BU major","BU school","BU start term","Story","Transcript lines","Units secured (matched)","Units pending (maybe/review)","No-credit lines"],
             [[s["student_id"], s["name"], s["intended_major"], s["bu_school"], s["bu_start_term"], s["story"], None, None, None, None] for s in D["students"]],
             {"Story": 70})

# Transcript lines with formula column
hdr = ["Line ID","Student ID","School","Term","Course code","Title","Credits","Calendar","Grade","Units if approved",
       "Expected status","Expected BU course","Expected units","Expected BU grade","Expected reason code","Flags","gt_counterpart","gt_overlap_pct"]
rows = []
for i, l in enumerate(D["lines"], 2):
    rows.append([f"L-{i-1:03d}", l["student_id"], l["school_code"], l["term"], l["course_code"], l["title"], l["credits"], l["calendar"], l["grade"],
                 f'=IF(H{i}="Quarter",ROUND(G{i}*2/3,2),G{i})', l["status"], l["bu_course"], l["units"], l["bu_grade"], l["reason_code"], l["flag"],
                 l["gt_counterpart"], l["gt_overlap_pct"]])
ws_l = sheet(wb, "Transcript Lines", hdr, rows, {"Title": 42, "Expected status": 48, "Flags": 60, "Units if approved": 12}, gt_cols=("gt_counterpart","gt_overlap_pct"))
n = len(rows) + 1
for r in range(2, len(D["students"]) + 2):
    ws_s[f"G{r}"] = f"=COUNTIF('Transcript Lines'!$B$2:$B${n},A{r})"
    ws_s[f"H{r}"] = f"=SUMIFS('Transcript Lines'!$M$2:$M${n},'Transcript Lines'!$B$2:$B${n},A{r},'Transcript Lines'!$K$2:$K${n},\"Matched*\")"
    ws_s[f"I{r}"] = f"=SUMIFS('Transcript Lines'!$J$2:$J${n},'Transcript Lines'!$B$2:$B${n},A{r},'Transcript Lines'!$K$2:$K${n},\"Maybe*\")+SUMIFS('Transcript Lines'!$J$2:$J${n},'Transcript Lines'!$B$2:$B${n},A{r},'Transcript Lines'!$K$2:$K${n},\"Pending*\")"
    ws_s[f"J{r}"] = f"=COUNTIFS('Transcript Lines'!$B$2:$B${n},A{r},'Transcript Lines'!$N$2:$N${n},\"TX\")"
    for c in "GHIJ": ws_s[f"{c}{r}"].font = BODY
ws_s.cell(len(D["students"]) + 3, 1, "Columns G-J are formulas over Transcript Lines. Units secured includes capped values; pending = units if approved.").font = Font(name="Arial", size=9, italic=True)

# BU courses
sheet(wb, "BU Courses", ["BU code","Dept","Title","Program","Requirement type","Level","Credits","Residency rule","Source","Description","Learning outcomes","Topics"],
      [[b["bu_code"], b["dept"], b["title"], b["program"], b["requirement_type"], b["level"], b["credits"], b["residency_rule"], b["source"],
        b["description"], J(b["learning_outcomes"]), J(b["topics"])] for b in D["bu"]],
      {"Description": 70, "Learning outcomes": 70, "Topics": 60, "Source": 40, "Residency rule": 40, "Title": 36})

# Schools
sheet(wb, "Schools", ["School code","Name","Type","Calendar","Accreditation","Location","Role","Note"],
      [[s["school_code"], s["name"], s["type"], s["calendar"], s["accreditation"], s["location"], s["role"], s["note"]] for s in D["schools"]], {"Note": 80, "Role": 45})

# External courses
sheet(wb, "External Courses", ["Ext ID","School","Course code","Title","Credits","Calendar","Units if transferred","Program area","Course type","Level","Source",
                               "Description","Learning outcomes","Topics","gt_counterpart","gt_overlap","gt_tier"],
      [[e["ext_id"], e["school_code"], e["course_code"], e["title"], e["credits"], e["calendar"], e["units_if_transferred"], e["program_area"],
        e["course_type"], e["level"], e["source"], e["description"], J(e["learning_outcomes"]), J(e["topics"]), e["gt_counterpart"], e["gt_overlap"], e["gt_tier"]]
       for e in D["ext"]],
      {"Description": 70, "Learning outcomes": 70, "Topics": 60, "Title": 40, "Source": 50}, gt_cols=("gt_counterpart","gt_overlap","gt_tier"))

# Decisions
sheet(wb, "Past Decisions", ["Decision ID","Ext ID","School","Source course","Source title","Source credits","Calendar","Decision","BU course","Units awarded",
                             "BU grade","Reason code","Overlap %","Decided by","Department","Decision date","Catalog year","Reusable","Split","Notes"],
      [[d["decision_id"], d["ext_id"], d["school_code"], d["source_course"], d["source_title"], d["source_credits"], d["calendar"], d["decision"],
        d["bu_course"], d["units_awarded"], ("TX" if d["decision"] == "No credit" else "T"), d["reason_code"], d["overlap_pct"], d["decided_by"],
        d["department"], d["decision_date"], d["catalog_year"], d["reusable"], d["split"], d["notes"]] for d in D["decisions"]],
      {"Source title": 40, "Notes": 45})

sheet(wb, "Syllabi Index", ["File","Ext ID","Course","School","State","Used by","Expected match (gt)","True overlap % (gt)","Note"],
      [[m["file"], m["ext_id"], m["course"], m["school"], m["state"], m["used_by"], m["expected_match"], m["true_overlap_pct"], m["note"]] for m in M["syllabi"]],
      {"Note": 70, "Course": 45, "File": 36}, gt_cols=("Expected match (gt)","True overlap % (gt)"))
sheet(wb, "Transcripts Index", ["File","Student ID","Student","School","Course lines"],
      [[m["file"], m["student_id"], m["student"], m["school"], m["courses"]] for m in M["transcripts"]], {"File": 50})
sheet(wb, "Rules", ["Rule ID","Category","Rule","Condition","Outcome","Reason code","Source"], RULES, {"Rule": 80, "Condition": 28, "Outcome": 36})
sheet(wb, "Reason Codes", ["Reason code","Decision","Meaning"], REASONS, {"Meaning": 60})
sheet(wb, "Test Cases", ["Test ID","Type","Input","Expected"], TESTS, {"Input": 70, "Expected": 70})

xlsx = os.path.join(OUT, "TransferIQ_Synthetic_Data.xlsx")
wb.save(xlsx)

# ---------------- seeds
sd = os.path.join(OUT, "seed"); os.makedirs(sd, exist_ok=True)
def wcsv(name, rows):
    if not rows: return
    with open(os.path.join(sd, name), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys())); w.writeheader()
        for r in rows: w.writerow({k: (J(v) if isinstance(v, list) else v) for k, v in r.items()})
wcsv("bu_courses.csv", D["bu"]); wcsv("schools.csv", D["schools"]); wcsv("external_courses.csv", D["ext"])
wcsv("past_decisions.csv", D["decisions"]); wcsv("students.csv", D["students"]); wcsv("transcript_lines_expected.csv", D["lines"])
wcsv("rules.csv", [dict(zip(["rule_id","category","rule","condition","outcome","reason_code","source"], r)) for r in RULES])
wcsv("reason_codes.csv", [dict(zip(["reason_code","decision","meaning"], r)) for r in REASONS])
wcsv("test_cases.csv", [dict(zip(["test_id","type","input","expected"], r)) for r in TESTS])
wcsv("syllabi_index.csv", M["syllabi"]); wcsv("transcripts_index.csv", M["transcripts"])
json.dump(dict(_note="Test data for the TransferIQ hackathon prototype. Source schools NEU/NYU use real catalog codes/titles; all students, grades, syllabi and decisions are invented and not issued by those schools. gt_* fields are ground truth: hide from the matcher.",
               bu_courses=D["bu"], schools=D["schools"], external_courses=D["ext"], past_decisions=D["decisions"], students=D["students"],
               transcript_lines_expected=D["lines"], rules=[dict(zip(["rule_id","category","rule","condition","outcome","reason_code","source"], r)) for r in RULES],
               reason_codes=[dict(zip(["reason_code","decision","meaning"], r)) for r in REASONS], files=M),
          open(os.path.join(sd, "transferiq_seed.json"), "w"), indent=1)
print("ok", xlsx)
